"""Markdown and DOCX rendering."""

from __future__ import annotations

import importlib
import importlib.util
import json
import os
import shutil
import stat
import subprocess
import sys
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
from conftest import FROZEN_NOW, make_transcript
from docx import Document
from docx.oxml.ns import qn
from docx.oxml.parser import parse_xml
from helpers_foundation import meeting as base_meeting
from helpers_foundation import minutes as base_minutes

from praktika.config import Settings
from praktika.eval import golden
from praktika.models import Classification, Decision, Flag, Ref, Review, Transcript
from praktika.render import docx as render_docx_mod
from praktika.render import markdown as md
from praktika.render.docx import render_docx
from praktika.render.markdown import BANNER_DRAFT, render_markdown

NOW = datetime.fromisoformat(FROZEN_NOW)
SECTIONS = [
    "## Attendees",
    "## Summary",
    "## Topics",
    "## Decisions",
    "## Actions",
    "## Open questions",
    "## Risks",
    "## Follow-ups",
    "## Reviewer flags",
    "## Provenance",
]


class Vault:
    def __init__(self, entries: dict[str, str]) -> None:
        self.entries = entries


def approved(**over: Any) -> Any:
    review = Review(
        status="approved",
        reviewer="f.khalid@acme.test",
        reviewer_source="session",
        reviewed_at=NOW,
    )
    return base_minutes(review=review, **over)


def seg_ref(t: Transcript, seg_id: str, quote: str | None = None) -> Ref:
    s = t.by_id()[seg_id]
    return Ref(
        segment_id=seg_id,
        start_s=s.start,
        end_s=s.end,
        speaker=s.speaker,
        quote=(quote if quote is not None else s.text)[:240],
    )


def test_markdown_sections_in_order(tmp_settings: Settings) -> None:
    t = make_transcript("en")
    text = render_markdown(base_minutes(), base_meeting(), t, settings=tmp_settings, now=NOW)
    positions = [text.index(h) for h in SECTIONS]
    assert positions == sorted(positions)
    assert text.startswith("**Classification: INTERNAL**")
    assert text.rstrip().endswith("**Classification: INTERNAL**")
    assert "# Minutes: Data team weekly" in text
    # times re-attached from the transcript, not from the ref
    assert "[S0005 00:00:40–00:00:49 F. Khalid]" in text
    assert "sha256 " + "a" * 64 in text and "Rendered at: " + NOW.isoformat() in text


def _assert_one_citation_per_line(text: str) -> int:
    """Every citation sits on a line of its own, and every gloss on the line after it; returns
    how many citation lines there are."""
    lines = text.splitlines()
    cited = [line for line in lines if "- [S0" in line]
    for line in cited:
        assert line.startswith("  - [S0") and line.count("[S0") == 1, line
    for line in lines:
        if "Gloss (EN):" in line:
            assert line.startswith("    - Gloss (EN): ") and "[S0" not in line, line
    return len(cited)


def _assert_tight(text: str, heading: str) -> None:
    """No blank line between the items of the list under ``heading`` (a blank line makes
    Markdown render the whole list loose)."""
    block: list[str] = []
    for line in text.split(heading + "\n", 1)[1].lstrip("\n").splitlines():
        if line and not line.startswith(("- ", "  ")):
            break
        block.append(line)
    while block and not block[-1]:
        block.pop()
    assert block and "" not in block, (heading, block)


def _cite_twice(m: Any) -> Any:
    """``m`` with every cited item citing its segments twice, so a run-together shows."""

    def doubled(items: list[Any]) -> list[Any]:
        return [i.model_copy(update={"refs": i.refs + i.refs}) if i.refs else i for i in items]

    fields = ["topics", "decisions", "actions", "open_questions", "risks", "flags"]
    fields += [
        f
        for f in ("matters_arising", "my_commitments", "their_commitments")
        if f in type(m).model_fields
    ]
    return m.model_copy(update={f: doubled(getattr(m, f)) for f in fields})


def _every_kind_of_item(kind: str, t: Transcript) -> Any:
    """Minutes of ``kind`` with every cited item type: several citations, an Arabic quote with
    its gloss, items without citations, matters arising and commitments."""
    from praktika.models import (
        MancomMinutes,
        MatterArising,
        OneToOneMinutes,
        OpenQuestion,
        Risk,
        TopicSummary,
    )

    ar = next(s.id for s in t.segments if s.language == "ar")
    base = base_minutes(
        topics=[
            TopicSummary(title="Pilot", summary="Agreed.",
                         key_points=["Volunteers only", "October"],
                         refs=[seg_ref(t, "S0001"), seg_ref(t, ar), seg_ref(t, "S0003")]),
            TopicSummary(title="Budget", summary="Noted.", key_points=[],
                         refs=[seg_ref(t, "S0002")]),
        ],
        decisions=[Decision(id="D1", statement="Volunteers only", kind="approved",
                            decided_by="F. Khalid", refs=[seg_ref(t, ar)])],
        open_questions=[
            OpenQuestion(id="Q1", question="Keep audio?", raised_by="F. Khalid", owner=None,
                         refs=[seg_ref(t, "S0004"), seg_ref(t, ar)]),
            OpenQuestion(id="Q2", question="Who signs?", raised_by=None, owner=None, refs=[]),
            OpenQuestion(id="Q3", question="When?", raised_by="L. Farouk", owner=None, refs=[]),
        ],
        risks=[
            Risk(id="R1", description="Names", severity="medium", owner=None,
                 mitigation="redact", refs=[seg_ref(t, ar)]),
            Risk(id="R2", description="Cost", severity="low", owner=None, mitigation=None,
                 refs=[]),
            Risk(id="R3", description="Time", severity="low", owner=None, mitigation="plan",
                 refs=[]),
        ],
        flags=[
            Flag(kind="name_to_verify", detail="Karim", priority=2, refs=[seg_ref(t, "S0002")]),
            Flag(kind="number_to_verify", detail="275,000", priority=2),
            Flag(kind="number_to_verify", detail="300,000", priority=2),
        ],
    )  # fmt: skip
    if kind == "general":
        return base
    data = base.model_dump()
    data["meeting_type"] = kind
    if kind == "mancom":
        data["matters_arising"] = [
            MatterArising(previous_action_id="A0", status="closed", note="Done",
                          refs=[seg_ref(t, ar), seg_ref(t, "S0001")]).model_dump(),
            MatterArising(previous_action_id="A9", status="open", note="Pending",
                          refs=[]).model_dump(),
            MatterArising(previous_action_id="A8", status="overdue", note="Late",
                          refs=[]).model_dump(),
        ]  # fmt: skip
        return MancomMinutes.model_validate(data)
    action = base.actions[0]
    data["my_commitments"] = [
        action.model_copy(update={"refs": [seg_ref(t, ar), seg_ref(t, "S0002")]}).model_dump()
    ]
    data["their_commitments"] = [
        action.model_copy(update={"id": "A2", "refs": [seg_ref(t, "S0003")]}).model_dump()
    ]
    return OneToOneMinutes.model_validate(data)


@pytest.mark.parametrize("kind", ["general", "mancom", "one_to_one"])
def test_every_citation_is_on_its_own_line(kind: str, tmp_settings: Settings) -> None:
    """In every template each citation, and each gloss, is on a line of its own (the templates
    run with trim_blocks, which once ran them together); an item without citations adds no
    blank line, so its list stays tight; and a topic's citations follow an 'Evidence:' line, so
    Markdown does not nest them under the topic's last key point."""
    t = make_transcript("mixed")
    m = _every_kind_of_item(kind, t)
    text = render_markdown(m, base_meeting(), t, settings=tmp_settings, now=NOW)
    assert _assert_one_citation_per_line(text) >= 10
    lines = text.splitlines()
    assert "    - Gloss (EN): Keep audio?" in lines and "    - Gloss (EN): Names" in lines
    assert "- **Q1** Keep audio? (raised by F. Khalid)" in lines
    assert "- **R1** [medium] Names — mitigation: redact" in lines
    for heading in ("## Open questions", "## Risks", "## Reviewer flags"):
        _assert_tight(text, heading)
    i = lines.index("- October")
    assert lines[i + 1 : i + 3] == ["", "Evidence:"] and lines[i + 3].startswith("  - [S0001 ")
    j = lines.index("Noted.")
    assert lines[j + 1] == "" and lines[j + 2].startswith("  - [S0002 "), "no key points"
    if kind == "mancom":
        _assert_tight(text, "## Matters arising")
        assert "- A0 — **closed**: Done" in lines and "    - Gloss (EN): Done" in lines
    if kind == "one_to_one":
        assert "    - Gloss (EN): Draft the notice" in lines


@pytest.mark.parametrize("gm", golden.load_all(), ids=lambda g: g.name)
def test_golden_minutes_put_every_citation_on_its_own_line(
    gm: golden.GoldenMeeting, tmp_settings: Settings
) -> None:
    """The same for the golden meetings, with every item, commitments included, citing twice."""
    m = _cite_twice(golden.minutes_from_playback(gm))
    text = render_markdown(m, gm.meeting, gm.transcript, settings=tmp_settings, now=NOW)
    assert _assert_one_citation_per_line(text) > 0


def test_draft_watermark_present_absent(tmp_settings: Settings) -> None:
    draft = render_markdown(base_minutes(), base_meeting(), None, settings=tmp_settings)
    assert draft.count(BANNER_DRAFT) == 2  # top banner and footer
    final = render_markdown(approved(), base_meeting(), None, settings=tmp_settings)
    assert BANNER_DRAFT not in final
    assert "Reviewer: f.khalid@acme.test (session)" in final
    assert "| Status | approved |" in final


def test_arabic_quote_verbatim_with_gloss(tmp_settings: Settings) -> None:
    t = make_transcript("mixed")
    ar_seg = next(s for s in t.segments if s.language == "ar")
    en_seg = next(s for s in t.segments if s.language == "en")
    m = base_minutes(
        decisions=[
            Decision(
                id="D1",
                statement="Pilot limited to volunteers",
                kind="approved",
                decided_by="F. Khalid",
                refs=[seg_ref(t, ar_seg.id)],
            ),
            Decision(
                id="D2",
                statement="Start on the first of October",
                kind="approved",
                decided_by="F. Khalid",
                refs=[seg_ref(t, en_seg.id)],
            ),
        ]
    )
    text = render_markdown(m, base_meeting(), t, settings=tmp_settings)
    assert f"“{ar_seg.text}”" in text
    assert "Gloss (EN): Pilot limited to volunteers" in text
    assert "Gloss (EN): Start on the first of October" not in text


def test_removed_items_listed_under_flags(tmp_settings: Settings) -> None:
    removed_text = "The vendor contract is signed next week (decided by Committee) — removed"
    m = base_minutes(
        flags=[
            Flag(kind="uncited_item_removed", detail=removed_text, item_json="{}", priority=1),
            Flag(kind="number_to_verify", detail="275,000 in: budget", priority=2),
        ]
    )
    text = render_markdown(m, base_meeting(), None, settings=tmp_settings)
    flags_at = text.index("## Reviewer flags")
    assert text.index("### Removed items") > flags_at
    assert text.index(removed_text) > flags_at
    assert removed_text not in text[:flags_at]
    assert "**number_to_verify** (priority 2): 275,000 in: budget" in text
    clean = render_markdown(base_minutes(), base_meeting(), None, settings=tmp_settings)
    assert "_No open flags._" in clean and "Removed items" not in clean


def test_docx_opens_and_has_banner_and_tables(tmp_settings: Settings, tmp_path: Path) -> None:
    t = make_transcript("mixed")
    ar_seg = next(s for s in t.segments if s.language == "ar")
    m = base_minutes(
        decisions=[
            Decision(
                id="D1",
                statement="Pilot approved",
                kind="approved",
                decided_by="F. Khalid",
                refs=[seg_ref(t, ar_seg.id)],
            )
        ],
        flags=[
            Flag(
                kind="uncited_item_removed",
                detail="Removed decision text",
                item_json="{}",
                priority=1,
            )
        ],
    )
    out = render_docx(m, base_meeting(), t, tmp_path / "x" / "minutes.docx", settings=tmp_settings)
    assert out.exists() and stat.S_IMODE(os.stat(out).st_mode) == 0o600
    doc = Document(str(out))
    header = doc.sections[0].header.paragraphs[0].text
    footer = doc.sections[0].footer.paragraphs[0].text
    assert "Classification: INTERNAL" in header and BANNER_DRAFT in header
    assert "Classification: INTERNAL" in footer and "Generator: qwen2.5:14b" in footer
    assert len(doc.tables) >= 3
    decisions = doc.tables[1]
    assert [c.text for c in decisions.rows[0].cells] == [
        "ID",
        "Decision",
        "Kind",
        "Decided by",
        "Evidence",
    ]
    assert decisions.rows[1].cells[1].text == "Pilot approved"
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "Removed decision text" in body and "Provenance" in body
    paragraphs = list(doc.paragraphs) + [
        p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs
    ]
    runs = [r for p in paragraphs for r in p.runs]
    rtl = [r for r in runs if r._r.rPr is not None and r._r.rPr.find(qn("w:rtl")) is not None]
    assert rtl and any(ar_seg.text in r.text for r in rtl)
    assert all(not md.ARABIC_RE.search(r.text) for r in runs if r not in rtl)


def test_docx_approved_has_no_watermark(tmp_settings: Settings, tmp_path: Path) -> None:
    out = render_docx(approved(), base_meeting(), None, tmp_path / "a.docx", settings=tmp_settings)
    doc = Document(str(out))
    assert BANNER_DRAFT not in doc.sections[0].header.paragraphs[0].text
    assert BANNER_DRAFT not in "\n".join(p.text for p in doc.paragraphs)


def test_detokenise_only_when_allowed(
    tmp_settings: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    vault = Vault({"«IBAN_1»": "BH67BMAG00001299123456", "«EMAIL_1»": "layla@example.test"})
    summary = "Payment to «IBAN_1» confirmed; notify «EMAIL_1»."
    draft = base_minutes(summary=summary)
    text = render_markdown(draft, base_meeting(), None, vault=vault, settings=tmp_settings)
    assert "«IBAN_1»" in text and "BH67BMAG00001299123456" not in text
    ok = render_markdown(
        approved(summary=summary), base_meeting(), None, vault=vault, settings=tmp_settings
    )
    assert "BH67BMAG00001299123456" in ok and "layla@example.test" in ok and "«" not in ok
    restricted = approved(summary=summary, classification=Classification.restricted)
    text = render_markdown(restricted, base_meeting(), None, vault=vault, settings=tmp_settings)
    assert "«IBAN_1»" in text and "BH67BMAG00001299123456" not in text
    blocked = approved(
        summary=summary, flags=[Flag(kind="identifier_detected", detail="iban", priority=1)]
    )
    text = render_markdown(blocked, base_meeting(), None, vault=vault, settings=tmp_settings)
    assert "«IBAN_1»" in text

    # Without a policy module the renderer fails closed.
    real_import = importlib.import_module

    def no_policy(name: str, *a: Any, **k: Any) -> Any:
        if name == "praktika.policy":
            raise ImportError(name)
        return real_import(name, *a, **k)

    monkeypatch.setattr(md.importlib, "import_module", no_policy)
    text = render_markdown(
        approved(summary=summary), base_meeting(), None, vault=vault, settings=tmp_settings
    )
    assert "«IBAN_1»" in text and "BH67BMAG00001299123456" not in text


@pytest.mark.parametrize("gm", golden.load_all(), ids=lambda g: g.name)
def test_golden_meetings_render_in_both_formats(
    gm: golden.GoldenMeeting, tmp_settings: Settings, tmp_path: Path
) -> None:
    m = golden.minutes_from_playback(gm)
    text = render_markdown(m, gm.meeting, gm.transcript, settings=tmp_settings)
    assert BANNER_DRAFT in text and "## Provenance" in text
    kind = gm.meeting.meeting_type.value
    if kind == "mancom":
        assert "# Management Committee minutes:" in text
        assert "## Figures mentioned" in text and "BHD 250,000" in text
        assert "## Escalations to the Board" in text
    elif kind == "one_to_one":
        assert "PRIVATE — not indexed" in text
        assert "### Organiser's commitments" in text and "### Other party's commitments" in text
        assert "Get Omar access to the GPU test server." in text.split("### Other party")[0]
    else:
        assert "# Minutes:" in text
    if gm.language != "en":
        assert "Gloss (EN):" in text
    out = render_docx(
        m, gm.meeting, gm.transcript, tmp_path / f"{gm.name}.docx", settings=tmp_settings
    )
    assert Document(str(out)).tables
    assert render_docx_mod.ARABIC_RE is md.ARABIC_RE


def test_docx_direction_follows_script_share(tmp_settings: Settings, tmp_path: Path) -> None:
    """A mostly English paragraph with one Arabic word stays LTR (no ``w:bidi``) and only the
    Arabic span is an RTL run; an Arabic paragraph is bidi; headings are handled too."""
    from praktika.render import docx as dx

    mixed = "“رانيا، you wanted to add the data platform migration?”"
    assert dx.paragraph_is_rtl(mixed) is False
    assert dx.paragraph_is_rtl("خلاص، نعتمد المرحلة الأولى بمبلغ 250,000 دينار") is True
    assert dx.paragraph_is_rtl("Budget: BHD 250,000") is False
    assert [a for _, a in dx.script_runs(mixed)] == [False, True, False]
    assert "".join(s for s, _ in dx.script_runs(mixed)) == mixed

    t = make_transcript("mixed")
    m = base_minutes(
        title="مراجعة سياسة الائتمان",
        summary=mixed,
        decisions=[
            Decision(
                id="D1",
                statement="Pilot approved",
                kind="approved",
                decided_by="F. Khalid",
                refs=[seg_ref(t, "S0001"), seg_ref(t, "S0005")],
            )
        ],
    )
    out = render_docx(m, base_meeting(), t, tmp_path / "dir.docx", settings=tmp_settings)
    doc = Document(str(out))

    def bidi(p: Any) -> bool:
        ppr = p._p.pPr
        return ppr is not None and ppr.find(qn("w:bidi")) is not None

    def rtl_runs(p: Any) -> list[str]:
        return [
            r.text
            for r in p.runs
            if r._r.rPr is not None and r._r.rPr.find(qn("w:rtl")) is not None
        ]

    summary = next(p for p in doc.paragraphs if mixed in p.text)
    assert not bidi(summary) and rtl_runs(summary) == ["رانيا،"]
    heading = next(p for p in doc.paragraphs if p.text == "مراجعة سياسة الائتمان")
    assert heading.style.name.startswith("Heading") and bidi(heading)
    assert rtl_runs(heading) == ["مراجعة سياسة الائتمان"]
    evidence = doc.tables[1].rows[1].cells[4]
    assert len(evidence.paragraphs) == 2, "one paragraph per citation, not one bidi block"
    assert not bidi(evidence.paragraphs[0]), "the English citation stays LTR"
    assert bidi(evidence.paragraphs[1]), "the Arabic citation is RTL on its own"
    assert rtl_runs(evidence.paragraphs[1]) == [t.by_id()["S0005"].text]


def test_markdown_wraps_arabic_quotes_in_bdi(tmp_settings: Settings) -> None:
    t = make_transcript("mixed")
    ar_seg = next(s for s in t.segments if s.language == "ar")
    en_seg = next(s for s in t.segments if s.language == "en")
    m = base_minutes(
        decisions=[
            Decision(
                id="D1",
                statement="x",
                kind="approved",
                decided_by="F. Khalid",
                refs=[seg_ref(t, ar_seg.id), seg_ref(t, en_seg.id)],
            )
        ]
    )
    text = render_markdown(m, base_meeting(), t, settings=tmp_settings)
    assert f"<bdi>“{ar_seg.text}”</bdi>" in text
    assert f"“{en_seg.text}”" in text and f"<bdi>“{en_seg.text}”" not in text


# --------------------------------------------------------------------------- DOCX metadata


def test_docx_core_properties_are_praktikas_not_the_template(
    tmp_settings: Settings, tmp_path: Path
) -> None:
    """Regression: every export carried python-docx's template metadata (author "python-docx",
    comments "generated by python-docx", created 2013)."""
    out = render_docx(
        base_minutes(), base_meeting(), None, tmp_path / "m.docx", settings=tmp_settings, now=NOW
    )
    cp = Document(str(out)).core_properties
    stamp = NOW.astimezone(UTC).replace(microsecond=0)
    assert cp.title == base_minutes().title == "Data team weekly"
    assert cp.author == "Praktika" and cp.last_modified_by == "Praktika"
    assert cp.created == stamp and cp.modified == stamp
    assert "python-docx" not in cp.comments
    assert cp.comments.startswith("Classification: INTERNAL") and BANNER_DRAFT in cp.comments
    assert cp.subject == "" and cp.keywords == "" and cp.category == ""
    assert cp.revision == 1


def test_docx_core_properties_default_to_export_time(
    tmp_settings: Settings, tmp_path: Path
) -> None:
    before = datetime.now(UTC).replace(microsecond=0)
    out = render_docx(approved(), base_meeting(), None, tmp_path / "a.docx", settings=tmp_settings)
    after = datetime.now(UTC)
    cp = Document(str(out)).core_properties
    assert before <= cp.created <= after and cp.modified == cp.created
    assert cp.comments == "Classification: INTERNAL", "no draft banner once approved"


_EXTENDED = "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}"


def _app_properties(docx_path: Path) -> dict[str, str]:
    """``docProps/app.xml`` of a saved document, element name -> text, after asserting that no
    part of the package names the template's Word build."""
    with zipfile.ZipFile(docx_path) as z:
        for name in z.namelist():
            assert b"Macintosh" not in z.read(name), name
        root = parse_xml(z.read("docProps/app.xml"))
        assert not any(n.startswith("docProps/thumbnail") for n in z.namelist())
    return {child.tag.removeprefix(_EXTENDED): child.text or "" for child in root}


def test_docx_app_properties_name_praktika_not_the_template(
    tmp_settings: Settings, tmp_path: Path
) -> None:
    """Regression: every export named the Word build python-docx's template was saved with as
    its application (the template's extended properties, copied verbatim)."""
    out = render_docx(approved(), base_meeting(), None, tmp_path / "a.docx", settings=tmp_settings)
    assert _app_properties(out) == {"Application": "Praktika"}
    assert Document(str(out)).paragraphs, "Word still opens it"


def test_md_to_docx_script_names_praktika_too(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    spec = importlib.util.spec_from_file_location(
        "md_to_docx", Path(__file__).resolve().parents[1] / "scripts" / "md_to_docx.py"
    )
    assert spec is not None and spec.loader is not None
    script = importlib.util.module_from_spec(spec)
    monkeypatch.setitem(sys.modules, "md_to_docx", script)  # dataclasses look the module up
    spec.loader.exec_module(script)
    src = tmp_path / "note.md"
    src.write_text("# Note\n\nOne paragraph.\n\n| a | b |\n|---|---|\n| 1 | 2 |\n", "utf-8")
    opts = script.Options("Calibri", "Consolas", 10.5, 9.0)
    out = script.convert(src, tmp_path / "note.docx", opts)
    assert _app_properties(out) == {"Application": "Praktika"}
    core = Document(str(out)).core_properties
    assert core.title == "Note" and core.author == "" and core.created.year >= 2026

    fenced = tmp_path / "fenced.md"
    fenced.write_text("```bash\n# [root] not a heading\n```\n\n# " + "Long " * 80 + "\n", "utf-8")
    title = Document(str(script.convert(fenced, tmp_path / "f.docx", opts))).core_properties.title
    assert title.startswith("Long Long") and len(title) <= 255, "fence skipped, title truncated"


def test_docx_core_title_is_one_line_within_the_opc_limit(
    tmp_settings: Settings, tmp_path: Path
) -> None:
    long = approved(title="Budget\nreview " + "x" * 400)
    out = render_docx(long, base_meeting(), None, tmp_path / "t.docx", settings=tmp_settings)
    title = Document(str(out)).core_properties.title
    assert title.startswith("Budget review x") and "\n" not in title and len(title) <= 255


# --------------------------------------------------------------------------- export contract


def _export_app(tmp_settings: Settings, tmp_path: Path, token: str) -> Any:
    from conftest import FakeIdentity

    from praktika import server
    from praktika.audit import AuditLog, JsonlAuditSink
    from praktika.models import MeetingState
    from praktika.store.repo import SqliteStore

    store = SqliteStore(tmp_path / "praktika.db")
    identity = FakeIdentity(source="session")
    audit = AuditLog(JsonlAuditSink(tmp_settings.data_dir / "audit.jsonl"), store, identity)
    store.save_meeting(base_meeting().model_copy(update={"state": MeetingState.approved}))
    store.save_transcript(make_transcript("en", n=12), delete_after=None)
    store.save_minutes(approved())
    return server.create_app(
        tmp_settings, store, identity, audit, clock=lambda: NOW, session_token=token
    )


@pytest.mark.parametrize("fmt", ["md", "docx"])
def test_export_route_needs_the_session_header_the_page_sends(
    tmp_settings: Settings, tmp_path: Path, fmt: str
) -> None:
    """The review page downloads exports with fetch() and the same headers as every other API
    call; a plain link sends none and is refused. Pin both halves of that contract."""
    from fastapi.testclient import TestClient

    from praktika import server

    token = server.new_session_token()
    app = _export_app(tmp_settings, tmp_path, token)
    mid = base_meeting().id
    link = TestClient(app, base_url="http://127.0.0.1")
    assert link.get(f"/api/export/{mid}.{fmt}").status_code == 401, "a bare link is refused"
    page = TestClient(
        app,
        base_url="http://127.0.0.1",
        headers={"X-Praktika-Review": "1", "X-Praktika-Token": token},
    )
    r = page.get(f"/api/export/{mid}.{fmt}")
    assert r.status_code == 200 and r.content
    disposition = r.headers["content-disposition"]
    assert disposition.startswith("attachment;") and f'filename="{mid}.v1.{fmt}"' in disposition


APP_JS = Path(render_docx_mod.__file__).resolve().parents[1] / "static" / "app.js"


def test_review_page_exports_with_fetch_not_a_bare_link() -> None:
    """Regression: the export buttons were ``<a href download>`` links, which cannot carry the
    session header, so every export got a 401. The page must fetch with ``apiHeaders``."""
    js = APP_JS.read_text(encoding="utf-8")
    assert 'export-md").href' not in js and 'export-docx").href' not in js
    body = js[js.index("async function downloadExport") :]
    body = body[: body.index("\n}\n")]
    assert "fetch(" in body and "apiHeaders(" in body
    assert "Content-Disposition" in body and "createObjectURL" in body
    assert "revokeObjectURL" in body and "notice(" in body


@pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")
def test_export_file_name_never_carries_a_path_separator() -> None:
    """Regression: the RFC 5987 ``filename*`` form returned the decoded name unsanitised, so
    ``..%2F..%2Fx`` came back as ``../../x``. Runs the real ``dispositionName`` under node."""
    js = APP_JS.read_text(encoding="utf-8")
    fn = js[js.index("function dispositionName") :]
    fn = fn[: fn.index("\n}\n") + 3]
    headers = [
        "attachment; filename*=UTF-8''..%2F..%2Fevil.md",
        'attachment; filename="a\\b.docx"',
        "attachment; filename=a/b.docx",
        'attachment; filename="m1.v1.md"',
        "attachment; filename*=UTF-8''%E0%A4%A",
    ]
    call = f"{json.dumps(headers)}.map(h => dispositionName(h, 'f.md'))"
    script = fn + f"console.log(JSON.stringify({call}));"
    out = subprocess.run(  # noqa: S603 - fixed local interpreter, script built from the repo
        [shutil.which("node") or "node", "-e", script],
        capture_output=True, text=True, timeout=20, check=True,
    )  # fmt: skip
    names = json.loads(out.stdout)
    assert names == [".._.._evil.md", "a_b.docx", "a_b.docx", "m1.v1.md", "f.md"]
